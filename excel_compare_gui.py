# -*- coding: utf-8 -*-
"""
Excel 比对工具（带 UI）
功能：加载两个 Excel 文件（.xlsx/.xlsm），一键比对，输出新 Excel：
  工作表1/2 = 两份文件的原始内容（差异单元格着色：文件1=红色、文件2=蓝色）
  工作表3 = 比对差异明细
支持：多工作表时选择"仅比对第一个工作表"或"比对全部工作表（按顺序配对）"；
长文本参数做片段级着色（仅实质差异段落标色，标点/细微差异不标色）。
"""
import os
import re
import sys
import math
import queue
import difflib
import threading
import subprocess
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
from copy import copy
from datetime import datetime, timedelta

import openpyxl
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

RED = 'FFFF0000'    # 文件1（第一个文件）差异色
BLUE = 'FF0000FF'   # 文件2（第二个文件）差异色

# ==================== 核心比对逻辑 ====================

def norm(v):
    """规范化单元格值用于比较"""
    if v is None:
        return ''
    if isinstance(v, float):
        return round(v, 6)
    return str(v).strip()

def is_diff(v1, v2):
    n1, n2 = norm(v1), norm(v2)
    if isinstance(n1, float) and isinstance(n2, float):
        return abs(n1 - n2) > 1e-6
    return n1 != n2

def split_similar_paras(a_list, b_list, ratio_th=0.9):
    """把两个段落列表拆分为 (eq_a, diff_a, diff_b)：
    相似度>=阈值(仅标点/细微差异)的段落视为相同(eq)，其余为差异(diff)"""
    eq_a, diff_a, diff_b = [], [], []
    used = [False] * len(b_list)
    for pa in a_list:
        best, best_j = -1, -1
        for j, pb in enumerate(b_list):
            if used[j]:
                continue
            if not pa.strip() and not pb.strip():
                r = 1.0
            else:
                r = difflib.SequenceMatcher(None, pa, pb, autojunk=False).ratio()
            if r > best:
                best, best_j = r, j
        if best >= ratio_th:
            used[best_j] = True
            eq_a.append(pa)
        else:
            diff_a.append(pa)
    for j, pb in enumerate(b_list):
        if not used[j]:
            diff_b.append(pb)
    return eq_a, diff_a, diff_b

def paragraph_diff(a, b):
    """按段落(换行)比对，返回 (a_parts, b_parts)，每项为 (text, is_diff)"""
    al = str(a).split('\n')
    bl = str(b).split('\n')
    sm = difflib.SequenceMatcher(None, al, bl, autojunk=False)
    a_parts, b_parts = [], []
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag == 'equal':
            a_parts.append(('\n'.join(al[i1:i2]), False))
            b_parts.append(('\n'.join(bl[j1:j2]), False))
        elif tag == 'replace':
            eq_a, diff_a, diff_b = split_similar_paras(al[i1:i2], bl[j1:j2])
            if eq_a:
                a_parts.append(('\n'.join(eq_a), False))
            if diff_a:
                a_parts.append(('\n'.join(diff_a), True))
            if diff_b:
                b_parts.append(('\n'.join(diff_b), True))
        elif tag == 'delete':
            a_parts.append(('\n'.join(al[i1:i2]), True))
        elif tag == 'insert':
            b_parts.append(('\n'.join(bl[j1:j2]), True))
    return a_parts, b_parts

def diff_summary(a, b, max_items=4, max_len=55):
    """生成差异要点摘要（两侧各取 max_items 条）"""
    al = str(a).split('\n')
    bl = str(b).split('\n')
    sm = difflib.SequenceMatcher(None, al, bl, autojunk=False)
    a_lines, b_lines = [], []
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag == 'replace':
            _, diff_a, diff_b = split_similar_paras(al[i1:i2], bl[j1:j2])
            for seg in diff_a:
                s = seg.strip()
                if s:
                    a_lines.append('【文件1独有】' + (s[:max_len] + ('…' if len(s) > max_len else '')))
            for seg in diff_b:
                s = seg.strip()
                if s:
                    b_lines.append('【文件2独有】' + (s[:max_len] + ('…' if len(s) > max_len else '')))
        elif tag == 'delete':
            for seg in al[i1:i2]:
                s = seg.strip()
                if s:
                    a_lines.append('【文件1独有】' + (s[:max_len] + ('…' if len(s) > max_len else '')))
        elif tag == 'insert':
            for seg in bl[j1:j2]:
                s = seg.strip()
                if s:
                    b_lines.append('【文件2独有】' + (s[:max_len] + ('…' if len(s) > max_len else '')))
    out = a_lines[:max_items] + b_lines[:max_items]
    if not out:
        return '（文本内容不同，详见内容列）'
    return '\n'.join(out)

def row_sig(ws, r):
    """行签名：该行所有非空单元格的规范化文本拼接"""
    parts = []
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=r, column=c).value
        if v is not None and str(v).strip():
            parts.append(str(norm(v)))
    return '|'.join(parts)

def align_sequence(seq_a, seq_b):
    """基于 SequenceMatcher 的对齐，返回配对列表 [(idx_a, idx_b)]，0 表示无对应（1-based）"""
    sm = difflib.SequenceMatcher(None, seq_a, seq_b, autojunk=False)
    pairs = []
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag == 'equal':
            for k in range(i2 - i1):
                pairs.append((i1 + k + 1, j1 + k + 1))
        elif tag == 'replace':
            n = max(i2 - i1, j2 - j1)
            for k in range(n):
                pairs.append((i1 + k + 1 if k < i2 - i1 else 0,
                              j1 + k + 1 if k < j2 - j1 else 0))
        elif tag == 'delete':
            for k in range(i2 - i1):
                pairs.append((i1 + k + 1, 0))
        elif tag == 'insert':
            for k in range(j2 - j1):
                pairs.append((0, j1 + k + 1))
    return pairs

def sig_fields(s):
    return [x for x in s.split('|') if x != '']

def sig_ratio(s1, s2):
    """行签名相似度：字段级（按|拆分）SequenceMatcher ratio"""
    f1, f2 = sig_fields(s1), sig_fields(s2)
    if not f1 and not f2:
        return 1.0
    sm = difflib.SequenceMatcher(None, f1, f2, autojunk=False)
    return sm.ratio()

def align_rows(ws_a, ws_b):
    """
    行对齐：
      equal（签名完全相同）→ 直接配对
      replace → 字段级相似度贪心匹配（ratio>=0.4 视为对应行），其余为独有行
      delete/insert → 独有行
    返回 (row_pairs, extras_a, extras_b, sigs_a, sigs_b)
    """
    max_r = max(ws_a.max_row, ws_b.max_row)
    sigs_a = [row_sig(ws_a, r) for r in range(1, max_r + 1)]
    sigs_b = [row_sig(ws_b, r) for r in range(1, max_r + 1)]
    sm = difflib.SequenceMatcher(None, sigs_a, sigs_b, autojunk=False)
    pairs, extras_a, extras_b = [], [], []
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag == 'equal':
            for k in range(i2 - i1):
                pairs.append((i1 + k + 1, j1 + k + 1))
        elif tag == 'replace':
            a_rows = list(range(i1 + 1, i2 + 1))
            b_rows = list(range(j1 + 1, j2 + 1))
            used_b = [False] * len(b_rows)
            for ra in a_rows:
                best, best_j = -1, -1
                for j, rb in enumerate(b_rows):
                    if used_b[j]:
                        continue
                    r = sig_ratio(sigs_a[ra - 1], sigs_b[rb - 1])
                    if r > best:
                        best, best_j = r, j
                if best >= 0.4:
                    used_b[best_j] = True
                    pairs.append((ra, b_rows[best_j]))
                else:
                    extras_a.append(ra)
            for j, rb in enumerate(b_rows):
                if not used_b[j]:
                    extras_b.append(rb)
        elif tag == 'delete':
            for k in range(i2 - i1):
                extras_a.append(i1 + k + 1)
        elif tag == 'insert':
            for k in range(j2 - j1):
                extras_b.append(j1 + k + 1)
    return pairs, extras_a, extras_b, sigs_a, sigs_b

def align_cells_nonempty(vals_a, vals_b, val_th=0.35):
    """
    行内单元格对齐：只对非空值做匹配，空单元格不参与。
    策略：
      1) 同列号强配对（内容行列结构稳定时最准确；"一侧数字一侧文本"除外，
         因为那是抬头区列布局不同，如 46156 在列5 vs "询价日期"标签在列5）；
      2) 剩余值按文本相似度贪心配对（处理抬头区等列布局不同的场景）；
      3) 仍未配对的记为独有。
    返回配对列表 [(ca, cb)]，ca/cb 为列号，0 表示无对应
    """
    def is_num(v):
        return isinstance(v, (int, float))
    nva = {i + 1: str(norm(v)) for i, v in enumerate(vals_a) if norm(v) != ''}
    nvb = {i + 1: str(norm(v)) for i, v in enumerate(vals_b) if norm(v) != ''}
    cols_a, cols_b = sorted(nva), sorted(nvb)
    pairs = []
    matched_b = set()
    # 1) 同列强配对
    for ca in cols_a:
        if ca in nvb:
            ra = is_num(vals_a[ca - 1])
            rb = is_num(vals_b[ca - 1])
            if not (ra != rb):  # 数字vs文本的例外
                pairs.append((ca, ca))
                matched_b.add(ca)
    # 2) 剩余值按相似度贪心配对
    rem_a = [c for c in cols_a if (c, c) not in pairs]
    for ca in rem_a:
        best, best_cb = -1, 0
        for cb in cols_b:
            if cb in matched_b:
                continue
            r = difflib.SequenceMatcher(None, nva[ca], nvb[cb], autojunk=False).ratio()
            if r > best:
                best, best_cb = r, cb
        if best >= val_th:
            matched_b.add(best_cb)
            pairs.append((ca, best_cb))
        else:
            pairs.append((ca, 0))
    # 3) 对方剩余
    for cb in cols_b:
        if cb not in matched_b:
            pairs.append((0, cb))
    pairs.sort(key=lambda p: (p[0] or 999, p[1]))
    return pairs

def compare_sheets(ws_a, ws_b):
    """
    比对两个工作表，返回 (diffs, col_header_map)
    diffs: [{loc1, loc2, item, field, dtype, v1, v2, summary, row_only}]
    col_header_map: {(c_a, c_b): 表头文本}，用于 Sheet3 字段列命名
    """
    row_pairs, extras_a, extras_b, sigs_a, sigs_b = align_rows(ws_a, ws_b)
    maxc = max(ws_a.max_column, ws_b.max_column)

    # 找表头行对（第一个两侧字段数>=5 且含"序号/名称"的行对），提取列名映射
    col_header_map = {}
    header_done = False
    for ra, rb in row_pairs:
        if not ra or not rb or header_done:
            continue
        fa, fb = sig_fields(sigs_a[ra - 1]), sig_fields(sigs_b[rb - 1])
        if len(fa) >= 5 and len(fb) >= 5 and ('序号' in fa or '名称' in fa or '序号' in fb or '名称' in fb):
            vals_a = [ws_a.cell(row=ra, column=c).value for c in range(1, maxc + 1)]
            vals_b = [ws_b.cell(row=rb, column=c).value for c in range(1, maxc + 1)]
            for ca, cb in align_cells_nonempty(vals_a, vals_b):
                if ca and cb:
                    va = norm(ws_a.cell(row=ra, column=ca).value)
                    vb = norm(ws_b.cell(row=rb, column=cb).value)
                    col_header_map[(ca, cb)] = va or vb
            header_done = True

    def field_name(ca, cb):
        for (x, y), h in col_header_map.items():
            if x == ca and y == cb:
                return h
        return f'第{ca}列'

    def item_of(vals, row):
        if len(vals) > 1 and vals[1] is not None and norm(vals[1]):
            return str(vals[1])[:40]
        return f'第{row}行'

    diffs = []

    # 独有行（跳过空行）
    for ra in extras_a:
        if not sigs_a[ra - 1]:
            continue
        content = '\n'.join(str(ws_a.cell(row=ra, column=c).value)
                            for c in range(1, ws_a.max_column + 1)
                            if ws_a.cell(row=ra, column=c).value is not None and str(ws_a.cell(row=ra, column=c).value).strip())
        diffs.append({
            'loc1': f'R{ra}（整行）', 'loc2': '—',
            'item': item_of([ws_a.cell(row=ra, column=c).value for c in range(1, maxc + 1)], ra),
            'field': '整行', 'dtype': '仅文件1有',
            'v1': content[:500], 'v2': '',
            'summary': f'该行仅存在于文件1（第{ra}行），文件2无对应行',
            'row_only': (ra, 0)
        })
    for rb in extras_b:
        if not sigs_b[rb - 1]:
            continue
        content = '\n'.join(str(ws_b.cell(row=rb, column=c).value)
                            for c in range(1, ws_b.max_column + 1)
                            if ws_b.cell(row=rb, column=c).value is not None and str(ws_b.cell(row=rb, column=c).value).strip())
        diffs.append({
            'loc1': '—', 'loc2': f'R{rb}（整行）',
            'item': item_of([ws_b.cell(row=rb, column=c).value for c in range(1, maxc + 1)], rb),
            'field': '整行', 'dtype': '仅文件2有',
            'v1': '', 'v2': content[:500],
            'summary': f'该行仅存在于文件2（第{rb}行），文件1无对应行',
            'row_only': (0, rb)
        })

    # 配对行：跳过两侧签名完全相同（equal）的行对
    for ra, rb in row_pairs:
        if not ra or not rb:
            continue
        if sigs_a[ra - 1] == sigs_b[rb - 1]:
            continue
        vals_a = [ws_a.cell(row=ra, column=c).value for c in range(1, maxc + 1)]
        vals_b = [ws_b.cell(row=rb, column=c).value for c in range(1, maxc + 1)]
        item = item_of(vals_a, ra)
        for ca, cb in align_cells_nonempty(vals_a, vals_b):
            va = vals_a[ca - 1] if ca else None
            vb = vals_b[cb - 1] if cb else None
            if ca and cb:
                if not is_diff(va, vb):
                    continue
                fname = field_name(ca, cb)
                sv1, sv2 = norm(va), norm(vb)
                if len(str(sv1)) > 80 and len(str(sv2)) > 80:
                    # 两侧都是长文本（参数/技术指标类）→ 片段级差异
                    dtype = '文本差异'
                    summary = diff_summary(va, vb)
                else:
                    dtype = '内容差异'
                    summary = f'文件1[{sv1 or "空"}] vs 文件2[{sv2 or "空"}]'
                diffs.append({
                    'loc1': f'R{ra}C{ca}', 'loc2': f'R{rb}C{cb}',
                    'item': item, 'field': fname, 'dtype': dtype,
                    'v1': va, 'v2': vb, 'summary': summary,
                    'row_only': None
                })
            else:
                dtype = '仅文件1有' if ca else '仅文件2有'
                if ca:
                    sv = norm(va)
                    diffs.append({
                        'loc1': f'R{ra}C{ca}', 'loc2': '—',
                        'item': item, 'field': field_name(ca, cb or ca), 'dtype': dtype,
                        'v1': va, 'v2': '',
                        'summary': f'该单元格内容仅文件1有：{str(sv)[:100]}',
                        'row_only': None
                    })
                else:
                    sv = norm(vb)
                    diffs.append({
                        'loc1': '—', 'loc2': f'R{rb}C{cb}',
                        'item': item, 'field': field_name(ca or cb, cb), 'dtype': dtype,
                        'v1': '', 'v2': vb,
                        'summary': f'该单元格内容仅文件2有：{str(sv)[:100]}',
                        'row_only': None
                    })
    return diffs, col_header_map

# ==================== 输出构建 ====================

def copy_ws(src, dst):
    """完整复制工作表（值+样式+列宽+行高+合并单元格）"""
    for key, dim in src.column_dimensions.items():
        dst.column_dimensions[key].width = dim.width
    for key, dim in src.row_dimensions.items():
        dst.row_dimensions[key].height = dim.height
    for mc in src.merged_cells.ranges:
        dst.merge_cells(str(mc))
    try:
        dst.freeze_panes = src.freeze_panes
    except Exception:
        pass
    for row in src.iter_rows():
        for c in row:
            if isinstance(c, openpyxl.cell.cell.MergedCell):
                continue
            d = dst.cell(row=c.row, column=c.column)
            d.font = copy(c.font)
            d.border = copy(c.border)
            d.fill = copy(c.fill)
            d.number_format = copy(c.number_format)
            d.protection = copy(c.protection)
            d.alignment = copy(c.alignment)
            d.value = c.value
    return dst

def colorize_cell(cell, color):
    """整格字体着色（保留其他样式）"""
    f = cell.font
    cell.font = Font(name=f.name, size=f.size, bold=f.bold, italic=f.italic,
                     underline=f.underline, strike=f.strike, color=color)

def rich_colorize(cell, text, other_text, color):
    """富文本：只把与对方不同的段落着色，相同段落保持原色"""
    f = cell.font
    if text is None:
        text = ''
    sv = str(text)
    if other_text is None or len(sv) < 60 or '\n' not in sv:
        colorize_cell(cell, color)
        return
    a_parts, b_parts = paragraph_diff(sv, str(other_text))
    blocks = []
    for t, isd in a_parts:
        if t == '':
            continue
        if isd:
            blocks.append(TextBlock(InlineFont(rFont=f.name, sz=f.size or 11, b=bool(f.bold), color=color), t))
        else:
            blocks.append(TextBlock(InlineFont(rFont=f.name, sz=f.size or 11, b=bool(f.bold)), t))
    cell.value = CellRichText(blocks)

def apply_diff_to_sheet(ws, loc, color, v, other_v):
    """把差异标到 ws 中 loc 指定的单元格"""
    if loc == '—' or '（整行）' in loc:
        return
    m = re.match(r'R(\d+)C(\d+)', loc)
    if not m:
        return
    r, c = int(m.group(1)), int(m.group(2))
    cell = ws.cell(row=r, column=c)
    rich_colorize(cell, v, other_v, color)

def build_output(path, pairs, extras, mode):
    """
    构建输出工作簿并保存
    pairs: [(sheet_a, ws_a, sheet_b, ws_b, diffs, col_map)] 逐个sheet对
    extras: (extra_sheets_a, extra_sheets_b) 仅某文件有的sheet
    mode: 'first' | 'all'
    返回 (保存路径, 各sheet差异数统计)
    """
    out = openpyxl.Workbook()
    out.remove(out.active)
    stats = []

    used_names = set()
    def unique_name(name):
        base = re.sub(r'[\[\]:*?/\\]', '_', str(name))[:28]
        n = base
        i = 1
        while n in used_names:
            i += 1
            n = f'{base}({i})'
        used_names.add(n)
        return n

    for k, (sa, ws_a, sb, ws_b, diffs, col_map) in enumerate(pairs):
        if mode == 'first' or len(pairs) == 1:
            if sa == sb:
                # 两个文件的工作表同名时加前缀区分
                n_a = unique_name(f'文件1_{sa}')
                n_b = unique_name(f'文件2_{sb}')
            else:
                n_a, n_b = unique_name(sa), unique_name(sb)
            n_d = unique_name('比对差异明细')
        else:
            n_a = unique_name(f'文件1_{sa}')
            n_b = unique_name(f'文件2_{sb}')
            n_d = unique_name(f'比对差异明细_{sa}')

        s1 = out.create_sheet(n_a)
        copy_ws(ws_a, s1)
        s2 = out.create_sheet(n_b)
        copy_ws(ws_b, s2)

        for d in diffs:
            apply_diff_to_sheet(s1, d['loc1'], RED, d['v1'], d['v2'])
            apply_diff_to_sheet(s2, d['loc2'], BLUE, d['v2'], d['v1'])
            if d.get('row_only'):
                ro, _ = d['row_only']
                if ro:
                    for c in range(1, ws_a.max_column + 1):
                        if s1.cell(row=ro, column=c).value is not None:
                            colorize_cell(s1.cell(row=ro, column=c), RED)
                else:
                    _, ro = d['row_only']
                    for c in range(1, ws_b.max_column + 1):
                        if s2.cell(row=ro, column=c).value is not None:
                            colorize_cell(s2.cell(row=ro, column=c), BLUE)

        # 差异明细工作表
        build_diff_sheet(out, n_d, diffs, sa, sb)
        stats.append((sa, sb, len(diffs)))

    # 仅某文件有的额外sheet：原样复制，不做比对
    extra_a, extra_b = extras
    for name, ws in extra_a:
        n = unique_name(f'文件1_{name}（仅文件1有）')
        copy_ws(ws, out.create_sheet(n))
    for name, ws in extra_b:
        n = unique_name(f'文件2_{name}（仅文件2有）')
        copy_ws(ws, out.create_sheet(n))

    out.save(path)
    return path, stats

def build_diff_sheet(out, title, diffs, sa, sb):
    """生成差异明细工作表"""
    ws = out.create_sheet(title)
    thin = Side(style='thin', color='FFBFBFBF')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    hdr_fill = PatternFill('solid', fgColor='FFD9E1F2')

    ws.merge_cells('A1:I1')
    ws['A1'] = f'比对差异明细（{sa} vs {sb}）'
    ws['A1'].font = Font(name='微软雅黑', size=14, bold=True)
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 26

    ws.merge_cells('A2:I2')
    ws['A2'] = ('图例：红色字体 = 文件1《' + sa + '》中与文件2不同/独有的内容；'
                '蓝色字体 = 文件2《' + sb + '》中与文件1不同/独有的内容。')
    ws['A2'].font = Font(name='微软雅黑', size=10, color='FF595959')
    ws['A2'].alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
    ws.row_dimensions[2].height = 30

    stat_types = {}
    for d in diffs:
        stat_types[d['dtype']] = stat_types.get(d['dtype'], 0) + 1
    stat_txt = '差异统计：共 ' + str(len(diffs)) + ' 处 ｜ ' + ('；'.join(f'{k} {v} 处' for k, v in stat_types.items()) if stat_types else '无差异')
    ws.merge_cells('A3:I3')
    ws['A3'] = stat_txt
    ws['A3'].font = Font(name='微软雅黑', size=10, bold=True, color='FFC00000')
    ws['A3'].alignment = Alignment(horizontal='left', vertical='center')
    ws.row_dimensions[3].height = 20

    headers = ['#', '位置（文件1）', '位置（文件2）', '项目/设备', '字段', '差异类型', '文件1内容', '文件2内容', '差异要点']
    cols = ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I']
    for ci, h in enumerate(headers):
        cell = ws[f'{cols[ci]}4']
        cell.value = h
        cell.font = Font(name='微软雅黑', size=10, bold=True)
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = border
    ws.row_dimensions[4].height = 22

    widths = [5, 13, 13, 20, 11, 11, 45, 45, 60]
    for ci, w in enumerate(widths):
        ws.column_dimensions[cols[ci]].width = w

    def est_height(text, width):
        if not text:
            return 16
        lines = 0
        for seg in str(text).split('\n'):
            lines += max(1, math.ceil(len(seg) / width))
        return min(160, max(16, lines * 13))

    r = 5
    for idx, d in enumerate(diffs, 1):
        v1s = '' if d['v1'] is None else str(d['v1'])
        v2s = '' if d['v2'] is None else str(d['v2'])
        vals = [idx, d['loc1'], d['loc2'], d['item'], d['field'], d['dtype'], v1s, v2s, d['summary']]
        for ci, val in enumerate(vals):
            cell = ws[f'{cols[ci]}{r}']
            cell.value = val
            cell.font = Font(name='微软雅黑', size=9)
            cell.alignment = Alignment(horizontal='left' if ci in (3, 5, 6, 7, 8) else 'center',
                                       vertical='top', wrap_text=True)
            cell.border = border
        tc = ws[f'F{r}']
        if d['dtype'] in ('仅文件1有', '仅文件2有'):
            tc.font = Font(name='微软雅黑', size=9, bold=True, color='FFC00000')
        elif d['dtype'] == '文本差异':
            tc.font = Font(name='微软雅黑', size=9, color='FF7030A0')
        else:
            tc.font = Font(name='微软雅黑', size=9, color='FF2E75B6')
        if d['dtype'] in ('文本差异', '内容差异', '仅文件1有', '仅文件2有'):
            if v1s:
                ws[f'G{r}'].font = Font(name='微软雅黑', size=9, color=RED)
            if v2s:
                ws[f'H{r}'].font = Font(name='微软雅黑', size=9, color=BLUE)
        ws.row_dimensions[r].height = max(est_height(v1s, 45), est_height(v2s, 45), est_height(d['summary'], 60), 18)
        r += 1

# ==================== GUI ====================

class ExcelCompareApp:
    def __init__(self, root):
        self.root = root
        root.title('Excel 比对工具')
        root.geometry('760x480')
        root.minsize(680, 420)

        self.path_a = tk.StringVar()
        self.path_b = tk.StringVar()
        self.path_out = tk.StringVar()
        self.log_q = queue.Queue()
        self.busy = False

        main = ttk.Frame(root, padding=12)
        main.pack(fill=tk.BOTH, expand=True)

        # 标题
        ttk.Label(main, text='Excel 文件比对工具', font=('微软雅黑', 14, 'bold')).pack(anchor=tk.W, pady=(0, 8))

        # 文件1
        row = ttk.Frame(main)
        row.pack(fill=tk.X, pady=3)
        ttk.Label(row, text='文件1：', width=7).pack(side=tk.LEFT)
        e1 = ttk.Entry(row, textvariable=self.path_a)
        e1.pack(side=tk.LEFT, fill=tk.X, expand=True)
        ttk.Button(row, text='浏览…', width=8, command=lambda: self.browse(1)).pack(side=tk.LEFT, padx=(6, 0))

        # 文件2
        row = ttk.Frame(main)
        row.pack(fill=tk.X, pady=3)
        ttk.Label(row, text='文件2：', width=7).pack(side=tk.LEFT)
        e2 = ttk.Entry(row, textvariable=self.path_b)
        e2.pack(side=tk.LEFT, fill=tk.X, expand=True)
        ttk.Button(row, text='浏览…', width=8, command=lambda: self.browse(2)).pack(side=tk.LEFT, padx=(6, 0))

        # 输出
        row = ttk.Frame(main)
        row.pack(fill=tk.X, pady=3)
        ttk.Label(row, text='输出：', width=7).pack(side=tk.LEFT)
        e3 = ttk.Entry(row, textvariable=self.path_out)
        e3.pack(side=tk.LEFT, fill=tk.X, expand=True)
        ttk.Button(row, text='浏览…', width=8, command=lambda: self.browse(3)).pack(side=tk.LEFT, padx=(6, 0))

        # 比对按钮
        btn_row = ttk.Frame(main)
        btn_row.pack(fill=tk.X, pady=(10, 4))
        self.btn_go = ttk.Button(btn_row, text='开始比对', command=self.start_compare)
        self.btn_go.pack(side=tk.LEFT)
        self.btn_open = ttk.Button(btn_row, text='打开结果文件', command=self.open_result, state=tk.DISABLED)
        self.btn_open.pack(side=tk.LEFT, padx=(8, 0))

        # 状态
        self.status = ttk.Label(main, text='就绪', foreground='#1a6b1a')
        self.status.pack(anchor=tk.W, pady=(4, 2))

        # 日志
        ttk.Label(main, text='运行日志：').pack(anchor=tk.W)
        self.log = tk.Text(main, height=12, font=('Consolas', 9), state=tk.DISABLED)
        self.log.pack(fill=tk.BOTH, expand=True, pady=(2, 0))
        self.last_result = None

        self.root.after(100, self._poll_log)

    # ---------- 界面操作 ----------
    def browse(self, which):
        if which == 3:
            p = filedialog.asksaveasfilename(
                title='选择输出文件', defaultextension='.xlsx',
                filetypes=[('Excel 文件', '*.xlsx'), ('所有文件', '*.*')],
                initialfile='比对结果.xlsx')
            if p:
                self.path_out.set(p)
            return
        p = filedialog.askopenfilename(
            title='选择 Excel 文件' + str(which),
            filetypes=[('Excel 文件', '*.xlsx *.xlsm'), ('所有文件', '*.*')])
        if p:
            if which == 1:
                self.path_a.set(p)
            else:
                self.path_b.set(p)
            self._auto_output()

    def _auto_output(self):
        pa = self.path_a.get()
        if pa and not self.path_out.get():
            d = os.path.dirname(pa)
            base = os.path.splitext(os.path.basename(pa))[0]
            self.path_out.set(os.path.join(d, base + '_比对结果.xlsx'))

    def log_line(self, msg):
        self.log_q.put(msg)

    def _poll_log(self):
        try:
            while True:
                msg = self.log_q.get_nowait()
                self.log.configure(state=tk.NORMAL)
                self.log.insert(tk.END, msg + '\n')
                self.log.see(tk.END)
                self.log.configure(state=tk.DISABLED)
        except queue.Empty:
            pass
        self.root.after(100, self._poll_log)

    def open_result(self):
        if self.last_result and os.path.exists(self.last_result):
            subprocess.Popen(['explorer.exe', '/select,', os.path.normpath(self.last_result)])
        else:
            messagebox.showinfo('提示', '还没有可打开的结果文件。')

    # ---------- 比对流程 ----------
    def start_compare(self):
        if self.busy:
            messagebox.showinfo('提示', '正在比对中，请稍候…')
            return
        pa, pb = self.path_a.get().strip(), self.path_b.get().strip()
        if not pa or not pb:
            messagebox.showwarning('提示', '请先选择两个 Excel 文件。')
            return
        if not os.path.isfile(pa):
            messagebox.showwarning('提示', f'文件不存在：{pa}')
            return
        if not os.path.isfile(pb):
            messagebox.showwarning('提示', f'文件不存在：{pb}')
            return
        out = self.path_out.get().strip()
        if not out:
            out = os.path.join(os.path.dirname(pa), os.path.splitext(os.path.basename(pa))[0] + '_比对结果.xlsx')
            self.path_out.set(out)
        try:
            wba = openpyxl.load_workbook(pa, read_only=True)
            names_a = list(wba.sheetnames)
            wba.close()
            wbb = openpyxl.load_workbook(pb, read_only=True)
            names_b = list(wbb.sheetnames)
            wbb.close()
        except Exception as e:
            messagebox.showerror('错误', f'无法读取 Excel 文件（请确认是 .xlsx/.xlsm 格式且未被占用）：\n{e}')
            return

        mode = 'first'
        if len(names_a) > 1 or len(names_b) > 1:
            mode = self.ask_sheet_mode(names_a, names_b)
            if mode is None:
                return  # 用户取消

        self.busy = True
        self.btn_go.configure(state=tk.DISABLED)
        self.btn_open.configure(state=tk.DISABLED)
        self.status.configure(text='比对中…', foreground='#b06000')
        t = threading.Thread(target=self._do_compare, args=(pa, pb, out, names_a, names_b, mode), daemon=True)
        t.start()

    def ask_sheet_mode(self, names_a, names_b):
        """弹窗询问比对范围，返回 'first'/'all'/None"""
        dlg = tk.Toplevel(self.root)
        dlg.title('选择比对范围')
        dlg.transient(self.root)
        dlg.grab_set()
        dlg.resizable(False, False)

        frm = ttk.Frame(dlg, padding=14)
        frm.pack()

        ttk.Label(frm, text='检测到文件包含多个工作表：', font=('微软雅黑', 10, 'bold')).grid(row=0, column=0, sticky=tk.W, pady=(0, 6))
        ttk.Label(frm, text=f'文件1（{os.path.basename(self.path_a.get())}）：{len(names_a)} 个工作表\n'
                            f'   ' + '、'.join(names_a), wraplength=420, justify=tk.LEFT).grid(row=1, column=0, sticky=tk.W)
        ttk.Label(frm, text=f'文件2（{os.path.basename(self.path_b.get())}）：{len(names_b)} 个工作表\n'
                            f'   ' + '、'.join(names_b), wraplength=420, justify=tk.LEFT).grid(row=2, column=0, sticky=tk.W, pady=(4, 8))

        var = tk.StringVar(value='first')
        ttk.Radiobutton(frm, text='仅比对第一个工作表', variable=var, value='first').grid(row=3, column=0, sticky=tk.W)
        ttk.Radiobutton(frm, text='比对全部工作表（按顺序一一配对）', variable=var, value='all').grid(row=4, column=0, sticky=tk.W, pady=(2, 10))

        result = {'v': None}

        def ok():
            result['v'] = var.get()
            dlg.destroy()

        def cancel():
            result['v'] = None
            dlg.destroy()

        btns = ttk.Frame(frm)
        btns.grid(row=5, column=0)
        ttk.Button(btns, text='确定', command=ok).pack(side=tk.LEFT, padx=6)
        ttk.Button(btns, text='取消', command=cancel).pack(side=tk.LEFT, padx=6)

        self.root.wait_window(dlg)
        return result['v']

    def _do_compare(self, pa, pb, out, names_a, names_b, mode):
        try:
            self.log_line(f'加载文件1：{pa}')
            wba = openpyxl.load_workbook(pa, data_only=True)
            self.log_line(f'加载文件2：{pb}')
            wbb = openpyxl.load_workbook(pb, data_only=True)

            if mode == 'first':
                pairs = [(names_a[0], wba[names_a[0]], names_b[0], wbb[names_b[0]], None, None)]
                extra_a, extra_b = [], []
            else:
                pairs = []
                n = min(len(names_a), len(names_b))
                for k in range(n):
                    pairs.append((names_a[k], wba[names_a[k]], names_b[k], wbb[names_b[k]], None, None))
                extra_a = [(nm, wba[nm]) for nm in names_a[n:]]
                extra_b = [(nm, wbb[nm]) for nm in names_b[n:]]

            # 逐个sheet对比对
            full_pairs = []
            total_diff = 0
            for k, (sa, ws_a, sb, ws_b, _, _) in enumerate(pairs):
                self.log_line(f'比对工作表 {sa} ↔ {sb} …')
                diffs, col_map = compare_sheets(ws_a, ws_b)
                total_diff += len(diffs)
                self.log_line(f'  {sa} vs {sb}：发现 {len(diffs)} 处差异')
                full_pairs.append((sa, ws_a, sb, ws_b, diffs, col_map))
            if extra_a:
                self.log_line(f'文件1 多余的 {len(extra_a)} 个工作表（仅复制）：' + '、'.join(nm for nm, _ in extra_a))
            if extra_b:
                self.log_line(f'文件2 多余的 {len(extra_b)} 个工作表（仅复制）：' + '、'.join(nm for nm, _ in extra_b))

            self.log_line(f'生成结果文件：{out}')
            path, stats = build_output(out, full_pairs, (extra_a, extra_b), mode)
            self.last_result = path
            self.log_line(f'完成！共发现 {total_diff} 处差异，结果已保存。')

            # 主线程刷新UI
            self.root.after(0, lambda: self._on_done(total_diff, stats))
        except Exception as e:
            import traceback
            err = traceback.format_exc()
            self.root.after(0, lambda: self._on_error(str(e), err))

    def _on_done(self, total_diff, stats):
        self.busy = False
        self.btn_go.configure(state=tk.NORMAL)
        self.btn_open.configure(state=tk.NORMAL)
        self.status.configure(text=f'完成：共 {total_diff} 处差异', foreground='#1a6b1a')
        detail = '\n'.join(f'  {sa} vs {sb}：{n} 处' for sa, sb, n in stats)
        messagebox.showinfo('比对完成',
                            f'比对完成，共发现 {total_diff} 处差异。\n\n{detail}\n\n结果已保存至：\n{self.last_result}')

    def _on_error(self, msg, tb):
        self.busy = False
        self.btn_go.configure(state=tk.NORMAL)
        self.status.configure(text='比对失败', foreground='#c00000')
        self.log_line('错误：' + msg)
        self.log_line(tb[-1500:])
        messagebox.showerror('错误', f'比对过程中出现错误：\n{msg}')


def main():
    # tkinter 在高 DPI 下清晰显示
    try:
        import ctypes
        ctypes.windll.shcore.SetProcessDpiAwareness(1)
    except Exception:
        pass
    root = tk.Tk()
    app = ExcelCompareApp(root)
    root.mainloop()


if __name__ == '__main__':
    main()
